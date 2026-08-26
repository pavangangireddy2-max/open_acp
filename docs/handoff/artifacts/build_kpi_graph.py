#!/usr/bin/env python3
# build_kpi_graph.py — generates kpi_graph.html ("KPI Flow Map" artifact 🕸️)
# Reads kra_training_sheet.xlsx (org tracker + 3 team views). Every edge comes from
# the curated LADDER map below (verbatim targets of each Section-B "Ladders to" cell)
# or from Section-A mirror formulas. Unmapped rows fail the build loudly — no silent drops.
import openpyxl, re, json, html as H, sys

XLSX = "kra_training_sheet.xlsx"
OUT  = "kpi_graph.html"

# ---------------- read org tracker ----------------
wb = openpyxl.load_workbook(XLSX, data_only=False)
org = wb["KPI Tracker FY26-27"]

FLAG_RE = re.compile(r"\s*\[(CASCADE|SAMPLE|CROSS-FUNCTION)\]")
def clean_metric(s):
    flags = FLAG_RE.findall(s or "")
    return re.sub(r"\s+", " ", FLAG_RE.sub("", s or "")).strip(), flags

kras = {}          # "KRA 1" -> {name, weight}
org_metrics = []   # collapsed metric groups
row2metric = {}    # sheet row -> metric index
cur_kra_key = None
cur = None
for r in range(3, org.max_row + 1):
    b = org.cell(r, 2).value
    if b:
        first = str(b).split("\n")[0].strip()
        mk = re.match(r"(KRA \d+)\s*[—-]\s*(.+)", first)
        if mk:
            cur_kra_key = mk.group(1)
            w = re.search(r"(\d+)\s*%", str(b)[len(first):])
            kras.setdefault(cur_kra_key, {"name": mk.group(2).strip(),
                                          "weight": (w.group(1) + "%") if w else ""})
    F, G = org.cell(r, 6).value, org.cell(r, 7).value
    if not G:            # note row
        cur = None
        continue
    if F:                # new metric group
        name, flags = clean_metric(str(F))
        cur = {"name": name, "flags": flags, "kra": cur_kra_key,
               "lane": str(org.cell(r, 12).value or "").strip(),
               "functions": str(org.cell(r, 11).value or "").strip(),
               "unit": str(org.cell(r, 13).value or "").strip(),
               "freq": str(org.cell(r, 14).value or "").strip(),
               "desc": str(org.cell(r, 8).value or "").strip(),
               "subrows": [], "rows": []}
        org_metrics.append(cur)
    assert cur is not None, f"org r{r}: KPI row with no open metric group"
    g = str(G)
    parts = g.split("::")
    head = parts[0].split(":") if len(parts) == 2 else []
    cohort = head[2] if len(head) >= 3 and head[2] else (head[1] if len(head) >= 2 and head[1] in ("All",) else "")
    label = parts[1].strip() if len(parts) == 2 else g
    budg = org.cell(r, 15).value
    cur["subrows"].append({"cohort": cohort, "label": label,
                           "budget": "" if budg is None else str(budg)})
    cur["rows"].append(r)
    row2metric[r] = len(org_metrics) - 1

assert len(org_metrics) == 20, f"expected 20 org metrics, got {len(org_metrics)}: {[m['name'] for m in org_metrics]}"
assert {"KRA 1", "KRA 2", "KRA 3", "KRA 4"} <= set(kras) <= {"KRA 1", "KRA 2", "KRA 3", "KRA 4", "KRA 5"}, sorted(kras)
# KRA 5 carries targets, not a weight; mapping deliberately not forced — hard-set it
kras["KRA 5"] = {"name": "Employability", "weight": "", "note": "B3 80% · B4 80% — mapping deliberately not forced"}

NORM = lambda s: re.sub(r"[–—]", "-", s or "").casefold().strip()
def org_id(key):
    hits = [i for i, m in enumerate(org_metrics) if NORM(key) in NORM(m["name"])]
    assert len(hits) == 1, f"org key {key!r} resolved to {[org_metrics[i]['name'] for i in hits]}"
    return f"o{hits[0]}"

# ---------------- read team views ----------------
TEAMS = [("fs",  "FS & CS Core",    "FullStack & CS Core View"),
         ("csi", "Content–CSI",     "CSI Team View"),
         ("cc",  "Content–Central", "Content–Central View")]
team_rows = {}   # tid -> list of B rows dicts
mirrors   = {}   # tid -> ordered unique metric ids
asks      = {}   # tid -> list of (category, name)
for tid, tlabel, sheet in TEAMS:
    ws = wb[sheet]
    hdr = next(r for r in range(1, 8)
               if "KPI name" in [str(ws.cell(r, c).value or "") for c in range(1, 16)])
    cols = {str(ws.cell(hdr, c).value or ""): c for c in range(1, 16)}
    cN, cS, cCat = cols["KPI name"], cols["Sect"], cols["Metric category"]
    cD, cL, cU, cF, cB = cols["Description"], cols["Ladders to"], cols["Unit"], cols["Freq"], cols["Budgeted"]
    B, mir, C = [], [], []
    for r in range(hdr + 1, ws.max_row + 1):
        name = ws.cell(r, cN).value
        if not name:
            continue
        sect = str(ws.cell(r, cS).value or "").strip()
        name = str(name)
        if name.startswith("="):
            mm = re.search(r"G(\d+)", name)
            if mm:
                mi = row2metric[int(mm.group(1))]
                if f"o{mi}" not in mir:
                    mir.append(f"o{mi}")
            continue
        if sect == "C":
            C.append((str(ws.cell(r, cCat).value or "").strip(), name.strip()))
            continue
        assert sect == "B", f"{sheet} r{r}: unexpected Sect={sect!r} for {name!r}"
        budg = ws.cell(r, cB).value
        B.append({"name": name.strip(),
                  "cat": str(ws.cell(r, cCat).value or "").strip(),
                  "desc": str(ws.cell(r, cD).value or "").strip(),
                  "ladder": str(ws.cell(r, cL).value or "").strip(),
                  "unit": str(ws.cell(r, cU).value or "").strip(),
                  "freq": str(ws.cell(r, cF).value or "").strip(),
                  "budget": "" if budg is None else str(budg)})
    team_rows[tid], mirrors[tid], asks[tid] = B, mir, C

assert [len(team_rows[t]) for t in ("fs", "csi", "cc")] == [29, 8, 8], [len(team_rows[t]) for t in ("fs","csi","cc")]
assert [len(mirrors[t])  for t in ("fs", "csi", "cc")] == [11, 7, 1], [len(mirrors[t]) for t in ("fs","csi","cc")]
assert [len(asks[t])     for t in ("fs", "csi", "cc")] == [4, 3, 2], [len(asks[t]) for t in ("fs","csi","cc")]

# ---------------- dept-layer nodes (§5 of the one-pager) ----------------
DEPT = {
 "bi":  ("Business Impact", "§5 block — Graded Assessment Achievement · Weekly Active Users",
         "The one-pager §5 Business Impact block. The team-owned numbers ARE the department scoreboard rows of the same name (Graded Assessment Achievement, Weekly Active Users)."),
 "cv":  ("Content Velocity", "§5 block — 4 delivery rows",
         "The §5 Content Velocity block: Learning Content Hours, Vernacular Content Hours, Practice & Assessment Pieces, Branding Content Assets Delivered. Each team row rolls into its matching §5 row."),
 "ce":  ("Content Efficiency", "§5 block — matching cost-per-item rows",
         "The §5 Content Efficiency block. Each team cost row rolls into its matching §5 row (cost per learning hour / vernacular hour / objective practice item / coding question / branding asset / runtime per learner / BOS approval; Central's spend-vs-plan reads the same block)."),
 "sa":  ("Stakeholder Alignment", "§5 block — request fulfillment · sprint delivery",
         "The §5 Stakeholder Alignment block: Stakeholder Content Request Fulfillment and Cross-functional Sprint Delivery. Central's Shared-Team Deliverables Landed rolls into the same block."),
 "pccc":("Product Capability Configuration Coverage", "§5 row · Product Learning Experience lane",
         "The §5 Program Delivery row “Product Capability Configuration Coverage” — share of journey steps where the required capability is live and configured. Domain Capability Delivery feeds it (plus the Learning Platform capability registry)."),
 "cru": ("Creative Resource Utilisation", "§5 row · PMO lane",
         "The §5 Executive Ops row “Creative Resource Utilisation” — designer and video-editor bandwidth used against planned deliverables. Central's All-Units roll-up feeds it; each team's own rate stays team-side."),
 "xru": ("Cross-functional Resource Utilisation", "§5 row · PMO lane",
         "The §5 Executive Ops row “Cross-functional Resource Utilisation” — allocated Engineering, Product, Pedagogy, DA/DE and Product Design bandwidth used. Central's All-Units roll-up feeds it."),
 "dtb": ("Develop the Best (mentoring)", "intent in §5 Executive Ops — not yet its own row",
         "“Developing the best (mentoring subordinates)” is named inside the §5 Executive Ops description. Power Performers Created (FS and CSI) is the measurable form of that intent; the §5 row split is deferred by agreement."),
}

# ---------------- curated ladder map ----------------
# key: (team, exact Section-B KPI name) -> list of targets or "HYG"
# targets: ("k", n) org KRA · ("o", org-metric key) tracker row · ("d", dept id) §5 · ("t", sibling row name)
L = {
 ("fs", "Summative Skill Assessment Achievement Rate"): [("o", "Formative → Summative Conversion"), ("k", 1)],
 ("fs", "Formative Skill Assessment Achievement Rate"): [("t", "Summative Skill Assessment Achievement Rate"), ("o", "Module-Quiz → Formative Conversion")],
 ("fs", "Graded Assessment Achievement Rate"): [("d", "bi")],
 ("fs", "Weekly Active Users"): [("d", "bi")],
 ("fs", "Pedagogy Initiative Impact"): [("o", "Practice → Module-Quiz Conversion"), ("k", 2)],
 ("fs", "Learner Accessed Content Completion Rate"): [("o", "Engagement-Matrix Cell Migration")],
 ("fs", "Practice Attempt-to-Completion Rate"): [("o", "Journey Step Health"), ("o", "Learning Environment Satisfaction")],
 ("fs", "Learning Content Hours Delivered"): [("d", "cv")],
 ("fs", "Vernacular Content Hours Delivered"): [("d", "cv")],
 ("fs", "Practice & Assessment Content Pieces Delivered"): [("d", "cv")],
 ("fs", "Branding Content Assets Delivered"): [("d", "cv")],
 ("fs", "Cost per Learning Hour Produced"): [("d", "ce")],
 ("fs", "Cost per Vernacular Content Hour"): [("d", "ce")],
 ("fs", "Cost per Objective Practice Item"): [("d", "ce")],
 ("fs", "Cost per Coding Question"): [("d", "ce")],
 ("fs", "Cost per Branding Content Asset"): [("d", "ce")],
 ("fs", "Platform Runtime Cost per Active Learner"): [("d", "ce")],
 ("fs", "Tech Stack Freshness Rate"): [("o", "Industry Update Adherence")],
 ("fs", "Stakeholder Content Request Fulfillment Rate"): [("d", "sa")],
 ("fs", "Cross-functional Sprint Delivery Rate"): [("d", "sa")],
 ("fs", "Evaluation Environment Coverage"): [("t", "Practice Attempt-to-Completion Rate"), ("o", "Journey Step Health"), ("o", "Learning Environment Satisfaction")],
 ("fs", "Domain Capability Delivery"): [("d", "pccc")],
 ("fs", "Cost of Operations"): "HYG",
 ("fs", "Roadmap Items Completion"): "HYG",
 ("fs", "Hires Made"): "HYG",
 ("fs", "Cost per Hire"): "HYG",
 ("fs", "Team Retention Rate"): "HYG",
 ("fs", "Power Performers Created"): [("d", "dtb")],
 ("fs", "Creative Resource Utilisation Rate"): "HYG",
 ("csi", "Cost per BOS Approval"): [("d", "ce")],
 ("csi", "Cost per Vernacular Content Hour"): [("d", "ce")],
 ("csi", "Cost of Operations"): "HYG",
 ("csi", "Roadmap Items Completion"): "HYG",
 ("csi", "Hires Made"): "HYG",
 ("csi", "Cost per Hire"): "HYG",
 ("csi", "Team Retention Rate"): "HYG",
 ("csi", "Power Performers Created"): [("d", "dtb")],
 ("cc", "Shared-Team Deliverables Landed"): [("d", "sa")],
 ("cc", "Shared-Team Spend vs Plan"): [("d", "ce")],
 ("cc", "Creative Resource Utilisation — All Units"): [("d", "cru")],
 ("cc", "Cross-functional Resource Utilisation — All Units"): [("d", "xru")],
 ("cc", "Shared Tool Adoption"): [("o", "Agentic Production Coverage")],
 ("cc", "Worklog & Status Hygiene — All Units"): "HYG",
 ("cc", "Check-ins Run"): "HYG",
 ("cc", "Actions Closed"): "HYG",
}

unmapped, used = [], set()
for tid, _, _ in TEAMS:
    for row in team_rows[tid]:
        if (tid, row["name"]) in L:
            used.add((tid, row["name"]))
        else:
            unmapped.append((tid, row["name"]))
stale = set(L) - used
if unmapped or stale:
    print("UNMAPPED Section-B rows:", unmapped)
    print("STALE map keys:", sorted(stale))
    sys.exit(1)

# extra org->KRA edge: Program Delivery Gap cascades into KRA 3 as well (tracker r29 note)
EXTRA_ORG_KRA = {"Program Delivery Gap": "KRA 3"}

# ---------------- build node & edge lists ----------------
TEAM_COLOR = {"fs": "#0D9488", "csi": "#B45309", "cc": "#6D28D9"}
nodes = {}   # id -> dict(kind, team, label, sub, sub2, x, y, w, h, tip)
edges = []   # (src, tgt, kind, team)   kind: lad | mir | kra

def tname_id(tid, name):
    for i, rowd in enumerate(team_rows[tid]):
        if rowd["name"] == name:
            return f"t_{tid}_{i}"
    raise AssertionError(f"sibling target {name!r} not in {tid}")

for tid, tlabel, _ in TEAMS:
    for i, rowd in enumerate(team_rows[tid]):
        tgt = L[(tid, rowd["name"])]
        if tgt == "HYG":
            continue
        for kind, val in tgt:
            dst = {"k": lambda v: f"k{v}", "o": org_id, "d": lambda v: f"d_{v}",
                   "t": lambda v: tname_id(tid, v)}[kind](val)
            edges.append((f"t_{tid}_{i}", dst, "lad", tid))
    for oid in mirrors[tid]:
        edges.append((f"a_{tid}", oid, "mir", tid))

for i, m in enumerate(org_metrics):
    edges.append((f"o{i}", "k" + m["kra"].split()[1], "kra", ""))
    if m["name"] in EXTRA_ORG_KRA:
        edges.append((f"o{i}", "k" + EXTRA_ORG_KRA[m["name"]].split()[1], "kra", ""))

# ---------------- layout ----------------
X_TEAM, W_TEAM = 14, 352
X_DEPT, W_DEPT = 566, 300
X_ORG,  W_ORG  = 992, 384
X_KRA,  W_KRA  = 1444, 164
WIDTH = 1620
NH_T, NH_D, NH_O, NH_K = 36, 46, 58, 66
PITCH_T, PITCH_O = 44, 70
CAT_H, GROUP_GAP, TOP = 24, 46, 118

y = TOP
team_meta = {}
for tid, tlabel, _ in TEAMS:
    gy = y
    nodes[f"a_{tid}"] = {"kind": "cap", "team": tid, "x": X_TEAM, "y": y, "w": W_TEAM, "h": 30,
                         "label": f"Section A — answers for {len(mirrors[tid])} org row{'s' if len(mirrors[tid]) != 1 else ''}",
                         "sub": "", "tip": ""}
    y += 30 + 12
    cat = None
    for i, rowd in enumerate(team_rows[tid]):
        if rowd["cat"] != cat:
            cat = rowd["cat"]
            nodes[f"c_{tid}_{i}"] = {"kind": "cat", "team": tid, "x": X_TEAM, "y": y,
                                     "w": W_TEAM, "h": 14, "label": cat, "sub": "", "tip": ""}
            y += CAT_H
        hyg = L[(tid, rowd["name"])] == "HYG"
        sub = " · ".join(x for x in (rowd["unit"], rowd["freq"]) if x)
        if hyg:
            sub = (sub + " · " if sub else "") + "hygiene"
        nodes[f"t_{tid}_{i}"] = {"kind": "team", "team": tid, "hyg": hyg,
                                 "x": X_TEAM, "y": y, "w": W_TEAM, "h": NH_T,
                                 "label": rowd["name"], "sub": sub, "row": rowd}
        y += PITCH_T
    team_meta[tid] = (gy, y)
    y += GROUP_GAP
HEIGHT_TEAMS = y

# org column: grouped by KRA in tracker order, centered on the team span
org_span = 4 * 30 + len(org_metrics) * PITCH_O
oy = max(TOP, (TOP + HEIGHT_TEAMS - GROUP_GAP - org_span) / 2)
seen_kra = None
for i, m in enumerate(org_metrics):
    if m["kra"] != seen_kra:
        seen_kra = m["kra"]
        nodes[f"g_{i}"] = {"kind": "kgrp", "team": "", "x": X_ORG, "y": oy, "w": W_ORG, "h": 14,
                           "label": f"{seen_kra} — {kras[seen_kra]['name']}", "sub": "", "tip": ""}
        oy += 30
    nodes[f"o{i}"] = {"kind": "org", "team": "", "x": X_ORG, "y": oy, "w": W_ORG, "h": NH_O,
                      "label": m["name"], "m": m,
                      "sub": m["lane"], "sub2": m["functions"]}
    oy += PITCH_O

# dept column: order by average source-y, sweep to remove overlaps
dsrc = {d: [] for d in DEPT}
for s, t, k, _ in edges:
    if t.startswith("d_"):
        dsrc[t[2:]].append(nodes[s]["y"])
order = sorted(DEPT, key=lambda d: sum(dsrc[d]) / len(dsrc[d]))
dy = None
for d in order:
    want = sum(dsrc[d]) / len(dsrc[d]) - NH_D / 2
    dy = want if dy is None else max(want, dy)
    dy = max(dy, TOP)
    nodes[f"d_{d}"] = {"kind": "dept", "team": "", "x": X_DEPT, "y": dy, "w": W_DEPT, "h": NH_D,
                       "label": DEPT[d][0], "sub": DEPT[d][1], "tipbody": DEPT[d][2]}
    dy += NH_D + 22

# KRA column: centroid of member org nodes, sweep; KRA 5 after KRA 4
ky = None
for n in (1, 2, 3, 4, 5):
    members = [nodes[s]["y"] + NH_O / 2 for s, t, k, _ in edges if t == f"k{n}"]
    want = (sum(members) / len(members) - NH_K / 2) if members else ky + NH_K + 60
    ky = want if ky is None else max(want, ky + NH_K + 24)
    kk = kras[f"KRA {n}"]
    nodes[f"k{n}"] = {"kind": "kra", "team": "", "x": X_KRA, "y": ky, "w": W_KRA, "h": NH_K,
                      "label": f"KRA {n}", "sub": kk["name"], "sub2": kk.get("weight", ""),
                      "note": kk.get("note", "")}
HEIGHT = int(max(HEIGHT_TEAMS, oy, dy, ky + NH_K) + 60)

# ---------------- tooltips ----------------
def esc(s): return H.escape(s, quote=True)
TLBL = dict((t[0], t[1]) for t in TEAMS)
tips = {}
for nid, n in nodes.items():
    if n["kind"] == "team":
        r = n["row"]
        head = f'<b>{esc(r["name"])}</b><span class="tteam t{n["team"]}">{esc(TLBL[n["team"]])}</span>'
        meta = " · ".join(x for x in (r["unit"], r["freq"], ("budget " + r["budget"]) if r["budget"] else "") if x)
        lad = f'<div class="tlad">Sheet says: {esc(r["ladder"])}</div>' if r["ladder"] else ""
        hygl = '<div class="thyg">Hygiene row — deliberately no org ladder; run by the team’s PM.</div>' if n.get("hyg") else ""
        tips[nid] = f'{head}<div class="tcat">{esc(r["cat"])} · Section B</div><div class="tdesc">{esc(r["desc"])}</div><div class="tmeta">{esc(meta)}</div>{lad}{hygl}'
    elif n["kind"] == "cap":
        names = " · ".join(org_metrics[int(o[1:])]["name"] for o in mirrors[n["team"]])
        tips[nid] = (f'<b>{esc(TLBL[n["team"]])} — Section A</b><div class="tdesc">Live mirrors of the org tracker rows this team answers for '
                     f'(same cells, not separate numbers):</div><div class="tmeta">{esc(names)}</div>')
    elif n["kind"] == "dept":
        tips[nid] = f'<b>{esc(n["label"])}</b><div class="tcat">Department scoreboard — one-pager §5</div><div class="tdesc">{esc(n["tipbody"])}</div>'
    elif n["kind"] == "org":
        m = n["m"]
        subs = "".join(f'<div class="tsub">{esc(s["cohort"] or "—")} · {esc(s["label"])}'
                       + (f' — budget {esc(s["budget"])}' if s["budget"] else " — budget TBD")
                       + "</div>" for s in m["subrows"])
        fl = (" " + " ".join(f'<span class="tflag">{f}</span>' for f in m["flags"])) if m["flags"] else ""
        meta = " · ".join(x for x in (m["unit"], m["freq"]) if x)
        tips[nid] = (f'<b>{esc(m["name"])}</b>{fl}<div class="tcat">Org tracker · {esc(m["kra"])} — {esc(kras[m["kra"]]["name"])}</div>'
                     f'<div class="tdesc">{esc(m["desc"])}</div>'
                     f'<div class="tmeta">Lane: {esc(m["lane"])}</div><div class="tmeta">Functions that move it: {esc(m["functions"])}</div>'
                     f'<div class="tmeta">{esc(meta)}</div>{subs}')
    elif n["kind"] == "kra":
        extra = f'<div class="tdesc">{esc(n["note"])}</div>' if n.get("note") else ""
        w = f' · weight {esc(n["sub2"])}' if n.get("sub2") else ""
        tips[nid] = f'<b>{esc(n["label"])} — {esc(n["sub"])}</b><div class="tcat">Org KRA{w}</div>{extra}'

# ---------------- svg ----------------
def trunc(s, w, px=6.35):
    lim = int(w / px)
    return s if len(s) <= lim else s[:lim - 1].rstrip() + "…"

S = []
S.append(f'<svg id="g" width="{WIDTH}" height="{HEIGHT}" viewBox="0 0 {WIDTH} {HEIGHT}" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="KPI flow map">')
# column headings
for x, w, t, sub in ((X_TEAM, W_TEAM, "TEAM KPIs — Section B (owned)", "FS & CS Core · Content–CSI · Content–Central"),
                     (X_DEPT, W_DEPT, "DEPT SCOREBOARD — §5", "one-pager blocks & rows"),
                     (X_ORG, W_ORG, "ORG TRACKER — FY26-27", "lane + functions on every row"),
                     (X_KRA, W_KRA, "ORG KRAs", "")):
    S.append(f'<text x="{x}" y="34" class="colh">{esc(t)}</text>')
    if sub: S.append(f'<text x="{x}" y="50" class="colsub">{esc(sub)}</text>')
# edges under nodes
def bez(s, t):
    a, b = nodes[s], nodes[t]
    x1, y1 = a["x"] + a["w"], a["y"] + a["h"] / 2
    x2, y2 = b["x"], b["y"] + b["h"] / 2
    dx = (x2 - x1) * 0.45
    return f"M{x1:.0f},{y1:.0f} C{x1+dx:.0f},{y1:.0f} {x2-dx:.0f},{y2:.0f} {x2:.0f},{y2:.0f}"
for i, (s, t, k, tm) in enumerate(edges):
    col = TEAM_COLOR.get(tm, "#9a9a94")
    dash = ' stroke-dasharray="6 5"' if k == "mir" else ""
    S.append(f'<path id="e{i}" class="edge e-{k}" d="{bez(s,t)}" stroke="{col}"{dash} data-s="{s}" data-t="{t}"/>')
# nodes
for nid, n in nodes.items():
    x, yy, w, h, k = n["x"], n["y"], n["w"], n["h"], n["kind"]
    if k == "cat":
        S.append(f'<text x="{x+2}" y="{yy+11}" class="catlab" data-team="{n["team"]}">{esc(n["label"].upper())}</text>')
        continue
    if k == "kgrp":
        S.append(f'<text x="{x+2}" y="{yy+11}" class="kgrplab">{esc(n["label"].upper())}</text>')
        continue
    cls = f'node nk-{k}' + (f' tm-{n["team"]}' if n["team"] else "") + (" hyg" if n.get("hyg") else "")
    S.append(f'<g class="{cls}" id="{nid}" data-team="{n.get("team","")}" tabindex="0">')
    S.append(f'<rect x="{x}" y="{yy}" width="{w}" height="{h}" rx="7"/>')
    if k == "cap":
        S.append(f'<text x="{x+12}" y="{yy+19}" class="ncap">{esc(n["label"])}</text>')
    elif k == "team":
        S.append(f'<text x="{x+12}" y="{yy+15}" class="nname">{esc(trunc(n["label"], w-22))}</text>')
        if n["sub"]:
            S.append(f'<text x="{x+12}" y="{yy+28}" class="nsub">{esc(trunc(n["sub"], w-22, 5.4))}</text>')
    elif k == "dept":
        S.append(f'<text x="{x+12}" y="{yy+18}" class="nname">{esc(trunc(n["label"], w-22))}</text>')
        S.append(f'<text x="{x+12}" y="{yy+33}" class="nsub">{esc(trunc(n["sub"], w-22, 5.4))}</text>')
    elif k == "org":
        S.append(f'<text x="{x+12}" y="{yy+17}" class="nname">{esc(trunc(n["label"], w-22))}</text>')
        S.append(f'<text x="{x+12}" y="{yy+32}" class="nsub">{esc(trunc("Lane: " + n["sub"], w-22, 5.4))}</text>')
        S.append(f'<text x="{x+12}" y="{yy+46}" class="nsub2">{esc(trunc("Moved by: " + n["sub2"], w-22, 5.1))}</text>')
    elif k == "kra":
        S.append(f'<text x="{x+12}" y="{yy+22}" class="kname">{esc(n["label"])}</text>')
        S.append(f'<text x="{x+12}" y="{yy+39}" class="ksub">{esc(trunc(n["sub"], w-20, 5.6))}</text>')
        if n.get("sub2"):
            S.append(f'<text x="{x+12}" y="{yy+54}" class="ksub2">weight {esc(n["sub2"])}</text>')
        elif n.get("note"):
            S.append(f'<text x="{x+12}" y="{yy+54}" class="ksub2">not forced — by design</text>')
    S.append("</g>")
S.append("</svg>")
SVG = "\n".join(S)

# ---------------- edge-list table ----------------
def node_name(nid):
    n = nodes[nid]
    pre = {"team": TLBL.get(n.get("team", ""), ""), "cap": TLBL.get(n.get("team", ""), "") + " Section A",
           "dept": "§5", "org": "Org", "kra": ""}.get(n["kind"], "")
    return (pre + " · " if pre and n["kind"] not in ("cap", "kra") else pre + (" — " if n["kind"] == "cap" else "")) + \
           (n["label"] + (" — " + n["sub"] if n["kind"] == "kra" else ""))
KINDL = {"lad": "ladders to", "mir": "answers for (mirror)", "kra": "rolls into"}
rows_html = "".join(f"<tr><td>{esc(node_name(s))}</td><td>{KINDL[k]}</td><td>{esc(node_name(t))}</td></tr>"
                    for s, t, k, _ in edges)
hyg_rows = [(TLBL[tid], r["name"]) for tid, _, _ in TEAMS for r in team_rows[tid] if L[(tid, r["name"])] == "HYG"]
hyg_html = " · ".join(f"{esc(t)}: {esc(nm)}" for t, nm in hyg_rows)
asks_html = " &nbsp;·&nbsp; ".join(
    f'<b>{esc(TLBL[tid])}</b>: ' + ", ".join(esc(nm) for _, nm in asks[tid]) for tid, _, _ in TEAMS)

# ---------------- html ----------------
EDGE_JSON = json.dumps([{"s": s, "t": t, "k": k, "m": tm} for s, t, k, tm in edges]).replace("</", "<\\/")
TIP_JSON = json.dumps(tips).replace("</", "<\\/")

html_out = f"""<meta charset="utf-8">
<title>KPI Flow Map</title>
<style>
  * {{ margin:0; padding:0; box-sizing:border-box; }}
  :root {{ --fs:#0D9488; --csi:#B45309; --cc:#6D28D9; --ink:#1a1a1a; --sub:#666; --line:#d8d8d3; --bg:#ffffff; }}
  body {{ font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif; background:var(--bg); color:var(--ink); padding:36px 20px 60px; line-height:1.5; }}
  .container {{ max-width:1660px; margin:0 auto; }}
  h1 {{ font-size:1.9em; font-weight:700; }}
  .subtitle {{ color:var(--sub); margin:4px 0 18px; max-width:72ch; }}
  .bar {{ display:flex; flex-wrap:wrap; gap:8px 14px; align-items:center; margin-bottom:6px; }}
  .chip {{ border:1px solid var(--line); border-radius:999px; padding:4px 12px; font-size:.82em; font-weight:600; cursor:pointer; background:#fff; color:var(--ink); }}
  .chip .dot {{ display:inline-block; width:9px; height:9px; border-radius:50%; margin-right:6px; vertical-align:1px; }}
  .chip.on {{ border-color:var(--ink); box-shadow:inset 0 0 0 1px var(--ink); }}
  .lg {{ font-size:.78em; color:var(--sub); display:flex; gap:16px; flex-wrap:wrap; align-items:center; margin-bottom:14px; }}
  .lg svg {{ vertical-align:middle; }}
  .hint {{ font-size:.82em; color:var(--sub); margin-bottom:10px; }}
  .wrap {{ overflow-x:auto; border:1px solid #eeeeea; border-radius:10px; background:#fff; }}
  svg#g {{ display:block; }}
  .colh {{ font-size:12.5px; font-weight:800; letter-spacing:.06em; fill:var(--ink); }}
  .colsub {{ font-size:10.5px; fill:var(--sub); }}
  .catlab {{ font-size:9.5px; font-weight:700; letter-spacing:.09em; fill:#8a8a84; }}
  .kgrplab {{ font-size:9.5px; font-weight:800; letter-spacing:.09em; fill:#8a8a84; }}
  .edge {{ fill:none; stroke-width:1.6; opacity:.38; transition:opacity .15s, stroke-width .15s; }}
  .node rect {{ fill:#fff; stroke:var(--line); stroke-width:1.2; transition:opacity .15s; cursor:pointer; }}
  .node text {{ pointer-events:none; }}
  .nk-team.tm-fs rect {{ stroke:var(--fs); }} .nk-team.tm-csi rect {{ stroke:var(--csi); }} .nk-team.tm-cc rect {{ stroke:var(--cc); }}
  .nk-team.hyg rect {{ stroke-dasharray:4 3; fill:#fcfcfa; }}
  .nk-cap rect {{ fill-opacity:.10; stroke:none; }}
  .nk-cap.tm-fs rect {{ fill:var(--fs); }} .nk-cap.tm-csi rect {{ fill:var(--csi); }} .nk-cap.tm-cc rect {{ fill:var(--cc); }}
  .nk-cap.tm-fs .ncap {{ fill:var(--fs); }} .nk-cap.tm-csi .ncap {{ fill:var(--csi); }} .nk-cap.tm-cc .ncap {{ fill:var(--cc); }}
  .ncap {{ font-size:11.5px; font-weight:700; }}
  .nk-dept rect {{ fill:#f5f5f2; stroke:#d2d2cc; }}
  .nk-org rect {{ fill:#fafaf8; }}
  .nk-kra rect {{ fill:var(--ink); stroke:var(--ink); }}
  .kname {{ fill:#fff; font-size:13px; font-weight:800; }}
  .ksub {{ fill:#fff; font-size:10.5px; font-weight:600; }}
  .ksub2 {{ fill:#bbb; font-size:9.5px; }}
  .nname {{ font-size:12px; font-weight:600; fill:var(--ink); }}
  .nsub {{ font-size:10px; fill:var(--sub); }}
  .nsub2 {{ font-size:9.5px; fill:#8a8a84; }}
  svg.sel .edge {{ opacity:.05; }}
  svg.sel .edge.hi {{ opacity:.95; stroke-width:2.4; }}
  svg.sel .node {{ opacity:.16; }}
  svg.sel .node.hi {{ opacity:1; }}
  svg.sel .node.hi rect {{ stroke-width:2; }}
  svg.tf .node[data-team]:not([data-team=""]):not(.tm-on) {{ opacity:.13; }}
  svg.tf .edge[data-m]:not(.m-on) {{ opacity:.04; }}
  svg.tf .node.dim {{ opacity:.13; }}
  #tip {{ position:fixed; z-index:9; max-width:380px; background:#fff; border:1px solid var(--ink); border-radius:8px; padding:10px 12px; font-size:.8em; box-shadow:0 6px 24px rgba(0,0,0,.14); display:none; pointer-events:none; }}
  #tip b {{ font-size:1.05em; }}
  .tteam {{ font-size:.85em; font-weight:700; margin-left:8px; padding:1px 7px; border-radius:999px; border:1px solid; }}
  .tfs {{ color:var(--fs); border-color:var(--fs); }} .tcsi {{ color:var(--csi); border-color:var(--csi); }} .tcc {{ color:var(--cc); border-color:var(--cc); }}
  .tcat {{ color:var(--sub); font-size:.9em; margin:2px 0 6px; }}
  .tdesc {{ margin-bottom:6px; }}
  .tmeta {{ color:var(--sub); font-size:.92em; }}
  .tsub {{ color:var(--sub); font-size:.92em; font-variant-numeric:tabular-nums; }}
  .tlad {{ margin-top:6px; padding-top:6px; border-top:1px dashed var(--line); font-style:italic; color:#444; }}
  .thyg {{ margin-top:6px; color:#8a6d00; }}
  .tflag {{ font-size:.78em; font-weight:700; border:1px solid var(--line); border-radius:4px; padding:0 4px; color:var(--sub); }}
  details {{ margin-top:22px; font-size:.86em; }}
  summary {{ cursor:pointer; font-weight:600; }}
  details table {{ border-collapse:collapse; margin-top:10px; width:100%; font-size:.95em; }}
  details td {{ border:1px solid var(--line); padding:5px 8px; }}
  .tscroll {{ overflow-x:auto; }}
  .foot {{ margin-top:22px; color:var(--sub); font-size:.8em; max-width:110ch; }}
  .foot p {{ margin-bottom:6px; }}
</style>
<div class="container">
  <h1>KPI Flow Map</h1>
  <p class="subtitle">How every KPI connects — from the rows each team owns (Section B), through the department scoreboard (one-pager §5), into the org tracker and the five KRAs. Hover or click any box to light up everything it touches; click a team chip to see just that team's flows.</p>
  <div class="bar" id="chips">
    <span class="chip on" data-t="">All teams</span>
    <span class="chip" data-t="fs"><span class="dot" style="background:var(--fs)"></span>FS &amp; CS Core</span>
    <span class="chip" data-t="csi"><span class="dot" style="background:var(--csi)"></span>Content–CSI</span>
    <span class="chip" data-t="cc"><span class="dot" style="background:var(--cc)"></span>Content–Central</span>
  </div>
  <div class="lg">
    <span><svg width="34" height="10"><line x1="0" y1="5" x2="34" y2="5" stroke="#555" stroke-width="2"/></svg> ladders to (owned number feeds it)</span>
    <span><svg width="34" height="10"><line x1="0" y1="5" x2="34" y2="5" stroke="#555" stroke-width="2" stroke-dasharray="6 4"/></svg> answers for (Section A mirror — same cell, not a separate number)</span>
    <span><svg width="34" height="10"><line x1="0" y1="5" x2="34" y2="5" stroke="#9a9a94" stroke-width="2"/></svg> rolls into its KRA</span>
    <span><svg width="30" height="14"><rect x="1" y="1" width="28" height="12" rx="4" fill="#fcfcfa" stroke="#999" stroke-dasharray="4 3"/></svg> hygiene — deliberately no org ladder (PM-run)</span>
  </div>
  <div class="wrap">{SVG}</div>
  <div id="tip"></div>
  <details><summary>Every connection as a list (for checking)</summary>
    <div class="tscroll"><table><tbody>{rows_html}</tbody></table></div>
  </details>
  <div class="foot">
    <p><b>Hygiene rows (no ladder, by design):</b> {hyg_html}.</p>
    <p><b>Section C asks are not drawn</b> — they are asks on other teams, not KPIs: {asks_html}.</p>
    <p><b>KRA 5 — Employability (B3 80% · B4 80%):</b> no KPI is forced into it by design. <b>KRA 4</b> also waits on GRIT-feature KPIs (pending definition with the GRIT owner). Dept-scoreboard boxes stop at §5 on purpose — that column is the department's own scoreboard; the org tracker column is the formal KRA path.</p>
    <p>Generated from kra_training_sheet.xlsx (org tracker + FS/CSI/Central team views) and the one-pager §5 · 22 Aug 2026.</p>
  </div>
</div>
<script>
const EDGES = {EDGE_JSON};
const TIPS = {TIP_JSON};
const svg = document.getElementById('g'), tip = document.getElementById('tip');
const down = {{}}, up = {{}};
EDGES.forEach((e, i) => {{ (down[e.s] = down[e.s] || []).push([e.t, i]); (up[e.t] = up[e.t] || []).push([e.s, i]); }});
function reach(id) {{
  const ns = new Set([id]), es = new Set();
  const walk = (start, adj) => {{ const q = [start], seen = new Set([start]);
    while (q.length) {{ const c = q.pop(); (adj[c] || []).forEach(([n, i]) => {{ es.add(i); ns.add(n); if (!seen.has(n)) {{ seen.add(n); q.push(n); }} }}); }} }};
  walk(id, down); walk(id, up);
  return [ns, es];
}}
let pinned = null;
function light(id) {{
  svg.classList.add('sel');
  svg.querySelectorAll('.hi').forEach(el => el.classList.remove('hi'));
  const [ns, es] = reach(id);
  ns.forEach(n => {{ const el = document.getElementById(n); if (el) el.classList.add('hi'); }});
  es.forEach(i => document.getElementById('e' + i).classList.add('hi'));
}}
function clearLight() {{ if (pinned) return; svg.classList.remove('sel'); svg.querySelectorAll('.hi').forEach(el => el.classList.remove('hi')); }}
svg.querySelectorAll('.node').forEach(g => {{
  g.addEventListener('mouseenter', () => {{ if (!pinned) light(g.id); showTip(g.id); }});
  g.addEventListener('mouseleave', () => {{ clearLight(); tip.style.display = 'none'; }});
  g.addEventListener('click', ev => {{ ev.stopPropagation(); pinned = (pinned === g.id) ? null : g.id; if (pinned) light(g.id); else clearLight(); }});
  g.addEventListener('focus', () => {{ if (!pinned) light(g.id); showTip(g.id, true); }});
  g.addEventListener('blur', () => {{ clearLight(); tip.style.display = 'none'; }});
}});
document.body.addEventListener('click', () => {{ pinned = null; svg.classList.remove('sel'); svg.querySelectorAll('.hi').forEach(el => el.classList.remove('hi')); }});
function showTip(id, focusMode) {{
  const h = TIPS[id]; if (!h) {{ tip.style.display = 'none'; return; }}
  tip.innerHTML = h; tip.style.display = 'block';
  if (focusMode) {{ const r = document.getElementById(id).getBoundingClientRect(); place(r.right, r.top); }}
}}
function place(x, y) {{
  const pad = 14, w = tip.offsetWidth, h = tip.offsetHeight;
  let L = x + pad, T = y + pad;
  if (L + w > innerWidth - 8) L = x - w - pad;
  if (T + h > innerHeight - 8) T = Math.max(8, innerHeight - h - 8);
  tip.style.left = Math.max(8, L) + 'px'; tip.style.top = T + 'px';
}}
svg.addEventListener('mousemove', ev => {{ if (tip.style.display === 'block') place(ev.clientX, ev.clientY); }});
// team filter chips
const chips = document.getElementById('chips');
const TEAMSETS = {{}};
['fs', 'csi', 'cc'].forEach(t => {{
  const ns = new Set(), es = new Set();
  svg.querySelectorAll('.node[data-team="' + t + '"]').forEach(g => {{ const [n2, e2] = reach(g.id); n2.forEach(x => ns.add(x)); e2.forEach(x => es.add(x)); }});
  TEAMSETS[t] = [ns, es];
}});
chips.addEventListener('click', ev => {{
  const c = ev.target.closest('.chip'); if (!c) return; ev.stopPropagation();
  chips.querySelectorAll('.chip').forEach(x => x.classList.remove('on')); c.classList.add('on');
  const t = c.dataset.t;
  svg.classList.remove('tf');
  svg.querySelectorAll('.tm-on, .m-on').forEach(el => el.classList.remove('tm-on', 'm-on'));
  svg.querySelectorAll('.node.dim').forEach(el => el.classList.remove('dim'));
  if (!t) return;
  svg.classList.add('tf');
  const [ns, es] = TEAMSETS[t];
  svg.querySelectorAll('.node').forEach(g => {{
    if (g.dataset.team === t) g.classList.add('tm-on');
    else if (!g.dataset.team && !ns.has(g.id)) g.classList.add('dim');
  }});
  EDGES.forEach((e, i) => {{ if (es.has(i) && (e.m === t || e.m === '')) document.getElementById('e' + i).classList.add('m-on'); }});
}});
</script>
"""
open(OUT, "w", encoding="utf-8").write(html_out)

n_nodes = sum(1 for n in nodes.values() if n["kind"] in ("team", "cap", "dept", "org", "kra"))
print(f"OK {OUT}: {len(html_out):,} bytes · {n_nodes} nodes ({sum(1 for n in nodes.values() if n['kind']=='team')} team / "
      f"{len(DEPT)} dept / {len(org_metrics)} org / 5 KRA / 3 capsules) · {len(edges)} edges "
      f"(lad {sum(1 for e in edges if e[2]=='lad')} · mir {sum(1 for e in edges if e[2]=='mir')} · kra {sum(1 for e in edges if e[2]=='kra')}) · "
      f"canvas {WIDTH}×{HEIGHT}")
assert html_out.count('class="edge') == len(edges)
assert "UNMAPPED" not in html_out
