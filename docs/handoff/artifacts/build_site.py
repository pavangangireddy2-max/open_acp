# Content OS — one site, four altitudes. Embeds the three published artifacts
# (HOD one-pager → KPI tracker → AI Engineer Ladder) with a site-generated Team
# level (FullStack & CS Core pilot, read from the master xlsx), so information
# flows org → department → team → individual on a single scroll.
#
# Technique: each artifact's CSS is wrapped in its section id via native CSS
# nesting (body{} hoisted out — all three share the identical family body).
# The site's own chrome (nav, altitude map, level bands, Level 3, download
# buttons) is authored here. The embedded ladder deliberately DIVERGES from the
# standalone artifact (site = wide audience; directives 2026-08-21): comp bands
# stripped and level 5 retitled "Head of [Domain Portfolio] Learning Systems"
# (directional, not finalized). Both master xlsx files are
# base64-embedded and offered through the viewer's `downloads` runtime
# capability (publish with capabilities={downloads: true}); the buttons stay
# hidden wherever the capability is absent — including local preview.
import base64
import json
import re
from openpyxl import load_workbook

ART = "/Users/pavan/Desktop/projects/open_acp/docs/handoff/artifacts/"
try:  # repo may be unreachable (e.g. macOS folder access revoked) — fall back to cwd copies
    open(ART + "hod_kpi_onepager.html", encoding="utf-8").close()
except OSError:
    ART = ""

def sub1(s, old, new):
    assert s.count(old) == 1, f"count={s.count(old)} for: {old[:70]!r}"
    return s.replace(old, new)

def load_artifact(fname):
    h = open(ART + fname, encoding="utf-8").read()
    css = re.search(r"<style[^>]*>(.*?)</style>", h, re.S).group(1)
    after = h.split("</style>", 1)[1]
    scripts = re.findall(r"<script[^>]*>.*?</script>", after, re.S)
    body = re.sub(r"<script[^>]*>.*?</script>", "", after, flags=re.S).strip()
    assert "<title>" not in body and "<style" not in body
    return css, body, scripts

FAMILY_BODY = ("font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; "
               "background: #ffffff; padding: 40px 20px; color: #1a1a1a; line-height: 1.5;")

def scope(css, sel):
    m = re.search(r"body\s*\{([^}]*)\}", css)
    assert m and " ".join(m.group(1).split()) == FAMILY_BODY, f"unexpected body in {sel}"
    assert "html" not in re.findall(r"(?:^|\})\s*([a-z]+)\s*\{", css)
    css = css.replace(m.group(0), "")
    return f"{sel} {{\n{css}\n}}"

op_css, op_body, op_scripts = load_artifact("hod_kpi_onepager.html")
tr_css, tr_body, tr_scripts = load_artifact("kra_training_sheet.html")
ld_css, ld_body, ld_scripts = load_artifact("role_cards.html")
assert not op_scripts and not ld_scripts and len(tr_scripts) == 1

# ---- ladder embed: site-only divergences from the standalone artifact ------
# ---- (comp stripped; level 5 retitled). Lineage sublines and the version
# ---- dialogue box left the source itself in the Aug-25 pass, so the strips
# ---- that used to remove them here are gone — the asserts below still guard.
ld_body = re.sub(r'<div class="lvl-comp">[^<|]*\|\s*([^<]*?)</div>',
                 r'<div class="lvl-comp">\1</div>', ld_body)
ld_body = sub1(ld_body,
    "Five levels, linear — an internship rung, then four employee rungs; AI Engineer 3 sits above Lead "
    "(org-wide scope) and the progression matrix below deliberately stops at Lead. Comp: Lead and AI Engineer 3 "
    "keep inherited bands; the changed rungs are with HR — no invented numbers. Strip the comp line before wide "
    "sharing if needed.",
    "Five levels, linear — an internship rung, then four employee rungs; above Lead sits the head-of-domain "
    "role (directional for now) and the progression matrix below deliberately stops at Lead. Comp bands are "
    "kept off this site — they live in the standalone ladder artifact and role_cards.xlsx.")
ld_body = sub1(ld_body,
    '<div class="t">AI Engineer 3 – [Domain Portfolio] Learning Systems</div>',
    '<div class="t">Head of [Domain Portfolio] Learning Systems</div>')
ld_body = sub1(ld_body,
    'Example: "AI Engineer 3 – Portfolio (All Domains) Learning Systems". Sets org-wide',
    '<em>Directional for now — title and shape indicative, not finalized.</em> '
    'Example: "Head of FullStack &amp; CS Core Learning Systems". Sets org-wide')
# v3 mentions the role in gates / stay bars / A-scale mapping / footnote too — retitle everywhere
ld_body = ld_body.replace("AI Engineer 3", "Head of [Domain Portfolio] Learning Systems")
ld_body = sub1(ld_body,
    "Defaults taken pending red-pen: comp bands included (artifact is private; strip for wide sharing) &middot; ",
    "Defaults taken pending red-pen: comp bands kept off this site (standalone artifact + xlsx carry them) &middot; ")
ld_body = sub1(ld_body,
    "level names Associate (internship) / Engineer / Senior / Lead / 3",
    "level names Associate (internship) / Engineer / Senior / Lead / Head-of-domain (directional)")
assert "L to " not in ld_body and "30L+" not in ld_body   # no comp band survives on the site
assert "formerly" not in ld_body and 'class="was"' not in ld_body
assert "AI Engineer 3" not in ld_body and "What changed in v3" not in ld_body

# ---- Level 3: FullStack & CS Core view, read from the master xlsx ----------
wbf = load_workbook(ART + "kra_training_sheet.xlsx")          # formulas (A refs)
ws = wbf["FullStack & CS Core View"]
trk = wbf["KPI Tracker FY26-27"]

def esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

A_ROWS = []                                                    # (T-no, kpi name)
for r in range(5, 20):
    tno = ws.cell(r, 1).value
    f = ws.cell(r, 6).value or ""
    m = re.search(r"G(\d+)", f)
    assert tno and m, (r, tno, f)
    A_ROWS.append((tno, trk.cell(int(m.group(1)), 7).value))
assert len(A_ROWS) == 15 and all(n for _, n in A_ROWS), A_ROWS

B_ROWS = [[ws.cell(r, c).value or "" for c in (1, 3, 4, 6, 7, 8, 9, 10, 14)]
          for r in range(21, 50)]
assert len(B_ROWS) == 29
C_ROWS = [[ws.cell(r, c).value or "" for c in (3, 6, 7)] for r in range(51, 55)]
assert len(C_ROWS) == 4
VIEW_TITLE = ws.cell(1, 1).value
LANE_LINE = ws.cell(2, 1).value

L3 = []
L3.append(f'<h1>{esc(VIEW_TITLE)}</h1>')
L3.append(f'<div class="subtitle">{esc(LANE_LINE)} &middot; pilot team view — the template every '
          'domain team copies &middot; live grid with budget/actual/variance columns: '
          '<strong>kra_training_sheet.xlsx</strong>, tab &ldquo;FullStack &amp; CS Core View&rdquo; '
          '(other live views there: CSI team — A 8 &middot; B 8 &middot; C 3 · Content&ndash;Central — '
          'A 1 &middot; B 8 &middot; C 2)</div>')
L3.append('<div class="key-point"><strong>Anatomy of a team view:</strong> '
          '<strong>Section A — inherited</strong>: the tracker rows above that this team works on through the '
          'Functions column; the number is answered for at department level. '
          '<strong>Section B — owned</strong>: the team&rsquo;s functional KPIs; each row names the tracker or '
          'department row it ladders to, so nothing here is an orphan metric. '
          '<strong>Section C — asks</strong>: what this team needs from other teams, named, so dependencies are '
          'contracts instead of surprises. Every view carries Product + Cohort columns — expansion slots for '
          'Academy / Intensive rows to land without a redesign.</div>')
L3.append(f'<h2>Section A — inherited from the tracker ({len(A_ROWS)} rows)</h2>')
L3.append('<div class="sectionlead">Mirrors of Level 2 rows (T-numbers) this team contributes to via '
          'Functions — shown as references here; the full rows with budgets live in the tracker above.</div>')
L3.append('<div class="achips">' + "".join(
    f'<span class="achip"><b>{esc(t)}</b> {esc(n)}</span>' for t, n in A_ROWS) + '</div>')
L3.append(f'<h2>Section B — owned by this team ({len(B_ROWS)} rows)</h2>')
L3.append('<div class="scroll"><table class="wide"><tr><th>#</th><th>Metric category</th><th>Product</th>'
          '<th>KPI name</th><th>Description</th><th>Ladders to</th><th>Unit</th><th>Freq</th><th>Remarks</th></tr>')
for sno, cat, prod, name, desc, ladder, unit, freq, rem in B_ROWS:
    prodc = f'<span class="prod">{esc(prod)}</span>' if prod != "All" else '<span class="muted">All</span>'
    L3.append(f'<tr><td class="muted">{esc(sno)}</td><td>{esc(cat)}</td><td>{prodc}</td>'
              f'<td class="kpiname">{esc(name)}</td><td>{esc(desc)}</td><td class="ladder">{esc(ladder)}</td>'
              f'<td>{esc(unit)}</td><td>{esc(freq)}</td><td class="muted">{esc(rem)}</td></tr>')
L3.append('</table></div>')
L3.append(f'<h2>Section C — asks of other teams ({len(C_ROWS)})</h2>')
L3.append('<div class="asks">')
for who, ask, detail in C_ROWS:
    L3.append(f'<div class="ask"><div class="ask-h"><b>{esc(ask)}</b><span class="of">of {esc(who)}</span></div>'
              f'<div class="ask-d">{esc(detail)}</div></div>')
L3.append('</div>')
tm_body = '<div class="container">\n' + "\n".join(L3) + '\n</div>'

tm_css = """  * { margin: 0; padding: 0; box-sizing: border-box; }
  .container { max-width: 1320px; margin: 0 auto; }
  h1 { font-size: 2em; margin-bottom: 8px; font-weight: 700; }
  .subtitle { font-size: 1em; color: #666; margin-bottom: 26px; max-width: 84ch; }
  h2 { font-size: 1.4em; margin: 38px 0 14px 0; padding-bottom: 8px; border-bottom: 2px solid #000; font-weight: 700; }
  .sectionlead { color: #666; font-size: .92em; margin: -4px 0 14px; max-width: 84ch; }
  .key-point { background: #e7f3ff; border-left: 3px solid #004085; padding: 12px 14px; margin: 18px 0; font-size: .93em; }
  .achips { display: flex; flex-wrap: wrap; gap: 6px; }
  .achip { border: 1px solid #999; border-radius: 3px; padding: 2px 8px; font-size: .82em; background: #fafafa; }
  .achip b { color: #004085; margin-right: 4px; }
  .scroll { overflow-x: auto; }
  table { width: 100%; border-collapse: collapse; margin-bottom: 26px; font-size: .84em; }
  table.wide { min-width: 1150px; }
  th { background: #000; color: #fff; padding: 8px 9px; text-align: left; font-weight: 600; border: 1px solid #000; }
  td { padding: 8px 9px; border: 1px solid #ccc; vertical-align: top; }
  .muted { color: #666; }
  .kpiname { font-family: ui-monospace, 'SF Mono', Menlo, monospace; font-size: .93em; font-weight: 600; }
  .ladder { color: #0b4f43; }
  .prod { border: 1px solid #999; border-radius: 3px; padding: 0 6px; font-size: .92em; white-space: nowrap; }
  .asks { display: grid; grid-template-columns: repeat(2, 1fr); gap: 12px; }
  .ask { border: 1px solid #ccc; border-radius: 4px; padding: 12px 14px; }
  .ask-h { display: flex; justify-content: space-between; gap: 10px; align-items: baseline; margin-bottom: 6px; }
  .ask-h .of { color: #666; font-size: .85em; white-space: nowrap; }
  .ask-d { font-size: .87em; color: #333; }
  @media (max-width: 860px) { .asks { grid-template-columns: 1fr; } }"""

# ---- downloadable masters: base64-embedded, saved via the viewer's
# ---- `downloads` capability. Buttons ship hidden; the script unhides them
# ---- only when claude.use("downloads") resolves (null → affordance absent).
XLSX_B64 = {f: base64.b64encode(open(ART + f, "rb").read()).decode("ascii")
            for f in ("kra_training_sheet.xlsx", "role_cards.xlsx")}

def dlbtn(name):
    return f'<button class="dlbtn" type="button" data-file="{name}" hidden>⬇ {name}</button>'

DL_SCRIPT = """<script>
(async () => {
  const XLSX = """ + json.dumps(XLSX_B64) + """;
  const dl = (window.claude && window.claude.use) ? await window.claude.use("downloads") : null;
  if (!dl) return;               // this view can't save files — buttons stay hidden
  document.querySelectorAll(".dlwrap").forEach(w => { w.hidden = false; });
  document.querySelectorAll(".dlbtn").forEach(b => {
    b.hidden = false;
    b.addEventListener("click", async () => {
      const name = b.dataset.file, label = b.textContent;
      b.disabled = true;
      try {
        await dl.save({ filename: name, data: Uint8Array.from(atob(XLSX[name]), c => c.charCodeAt(0)) });
        b.textContent = "saved \\u2713";
      } catch (e) {
        b.textContent = (e && e.code === "declined") ? label : "couldn't save \\u2014 try again";
      }
      setTimeout(() => { b.textContent = label; b.disabled = false; }, 1800);
    });
  });
})();
</script>"""

# ---- site chrome -----------------------------------------------------------
SITE_CSS = """  * { margin: 0; padding: 0; box-sizing: border-box; }
  body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
    background: #ffffff; color: #1a1a1a; line-height: 1.5; }
  #topnav { position: sticky; top: 0; z-index: 60; background: #ffffff; border-bottom: 2px solid #000;
    display: flex; align-items: center; gap: 18px; padding: 10px 20px; flex-wrap: wrap; }
  #topnav .brand { font-weight: 800; letter-spacing: .02em; }
  #topnav a { color: #1a1a1a; text-decoration: none; font-size: .9em; padding: 2px 2px; border-bottom: 2px solid transparent; }
  #topnav a:hover, #topnav a:focus-visible { border-bottom-color: #000; outline: none; }
  #topnav .lvl-no { color: #666; font-variant-numeric: tabular-nums; margin-right: 3px; }
  .site-section { scroll-margin-top: 58px; }
  .pad { padding: 40px 20px; }
  .hero { max-width: 1180px; margin: 0 auto; }
  .hero h1 { font-size: 2.4em; font-weight: 800; margin: 18px 0 10px; }
  .hero .tag { color: #666; font-size: 1.02em; max-width: 74ch; }
  .flow { max-width: 880px; margin: 34px auto 8px; }
  .fcard { display: block; color: inherit; text-decoration: none; border: 1px solid #ccc; border-radius: 5px;
    padding: 14px 16px; display: grid; grid-template-columns: 138px 1fr; gap: 14px; background: #fff; }
  .fcard:hover, .fcard:focus-visible { border-color: #000; outline: none; }
  .fcard .lv { font-weight: 800; font-size: .84em; letter-spacing: .06em; text-transform: uppercase; }
  .fcard .lv .n { display: block; font-size: 2em; line-height: 1.1; }
  .fcard .what { font-size: .92em; }
  .fcard .what b { display: block; margin-bottom: 2px; font-size: 1.05em; }
  .fcard .meta { color: #666; font-size: .88em; margin-top: 4px; }
  .conn { display: flex; justify-content: space-between; align-items: center; padding: 6px 18px; font-size: .82em; }
  .conn .down { color: #004085; }
  .conn .up { color: #0e6e5c; }
  .legendline { text-align: center; color: #666; font-size: .85em; margin-top: 14px; }
  .legendline .down { color: #004085; } .legendline .up { color: #0e6e5c; }
  .band { background: #000; color: #fff; padding: 26px 20px; }
  .band .in { max-width: 1180px; margin: 0 auto; display: flex; flex-wrap: wrap; gap: 8px 22px; align-items: baseline; }
  .band .lv { font-size: .8em; letter-spacing: .12em; text-transform: uppercase; color: #bbb; }
  .band .nm { font-size: 1.35em; font-weight: 800; }
  .band .desc { color: #ddd; font-size: .92em; flex-basis: 100%; max-width: 90ch; }
  .band .links { margin-left: auto; display: flex; gap: 10px; }
  .band a { color: #fff; font-size: .82em; text-decoration: none; border: 1px solid #666; border-radius: 3px; padding: 2px 9px; white-space: nowrap; }
  .band a:hover, .band a:focus-visible { border-color: #fff; outline: none; }
  .band .chipfile { color: #bbb; font-size: .82em; border: 1px dashed #555; border-radius: 3px; padding: 2px 9px; white-space: nowrap; }
  footer { border-top: 2px solid #000; margin-top: 20px; padding: 30px 20px 46px; }
  footer .in { max-width: 1180px; margin: 0 auto; font-size: .88em; color: #444; }
  footer h3 { font-size: 1em; margin-bottom: 8px; }
  footer ul { list-style: none; }
  footer li { margin: 3px 0; }
  footer .mono { font-family: ui-monospace, 'SF Mono', Menlo, monospace; font-size: .92em; }
  footer a { color: #004085; }
  .dlbtn { font: inherit; font-size: .82em; cursor: pointer; white-space: nowrap;
    border: 1px solid #0e6e5c; border-radius: 3px; background: #0e6e5c; color: #fff; padding: 2px 9px; }
  .dlbtn:hover, .dlbtn:focus-visible { background: #0b574a; border-color: #0b574a; outline: none; }
  .dlbtn[disabled] { opacity: .65; cursor: default; }
  @media (max-width: 700px) { .fcard { grid-template-columns: 1fr; gap: 6px; } }"""

LEVELS = [
    ("org", "1", "Organization", "🎯 HOD one-pager",
     "What the org holds this department to — the KRAs, and the department KPI baseline that answers them.",
     "Read by: Founders &middot; HOD. The contract everything below serves.",
     "https://claude.ai/code/artifact/cd52be82-3fd5-4264-a7aa-5b4e5e132ed9", None),
    ("dept", "2", "Department", "🧭 KPI tracker",
     "The 28 rows the department runs on — metric category says <em>what</em>, lane says <em>who answers</em>, cadence says <em>when</em>.",
     "Read by: HOD &middot; lane leads &middot; PMs. Master: kra_training_sheet.xlsx.",
     "https://claude.ai/code/artifact/44869dfd-506f-4f4c-9d2b-0581b8e065d1", "kra_training_sheet.xlsx"),
    ("team", "3", "Team", "Team scorecards",
     "One team's slice: Section A inherited from the tracker &middot; Section B owned &middot; Section C asks of other teams. Pilot below: FullStack &amp; CS Core.",
     "Read by: team leads &middot; SMEs. Views live as tabs in the tracker master xlsx.",
     None, "kra_training_sheet.xlsx"),
    ("ind", "4", "Individual", "🪜 AI Engineer Ladder",
     "An internship rung plus four employee levels, the team's Project Manager card, promotion gates and stay bars, 21 progression areas, and the rating rubric — every area wired to the same KPI rows the levels above run on.",
     "Read by: every engineer &middot; their manager. Master: role_cards.xlsx.",
     "https://claude.ai/code/artifact/dc3ab90d-c4a4-448d-bdb5-fb43a9235734", "role_cards.xlsx"),
]

def band(key, no, name, art, desc, meta, url, master):
    links = []
    if url:
        links.append(f'<a href="{url}" target="_blank" rel="noopener">standalone artifact ↗</a>')
    if master:
        links.append(f'<span class="chipfile">master: {master}</span>')
        links.append(dlbtn(master))
    return (f'<div class="band" id="{key}-band"><div class="in">'
            f'<span class="lv">Level {no} · {name}</span><span class="nm">{art}</span>'
            f'<span class="links">{"".join(links)}</span>'
            f'<span class="desc">{desc}</span></div></div>')

NAV = ('<nav id="topnav"><span class="brand">Content OS</span>'
       '<a href="#map">The map</a>'
       + "".join(f'<a href="#{k}"><span class="lvl-no">{n}</span>{nm}</a>'
                 for k, n, nm, *_ in LEVELS)
       + '</nav>')

HERO = ['<div class="pad site-section" id="map"><div class="hero">']
HERO.append('<h1>Content OS</h1>')
HERO.append('<div class="tag">The Content &amp; Curriculum department, end to end, on one page. Four altitudes: '
            'the organization sets the targets, the department runs the tracker, teams own their scorecards, and '
            'individual careers are wired to the very same rows. <strong>Targets cascade down; numbers ladder back '
            'up the same wires.</strong> Scroll, or jump to your altitude.</div>')
HERO.append('<div class="flow">')
for i, (key, no, name, art, desc, meta, url, master) in enumerate(LEVELS):
    if i:
        HERO.append('<div class="conn"><span class="down">▼ targets cascade</span>'
                    '<span class="up">numbers ladder up ▲</span></div>')
    HERO.append(f'<a class="fcard" href="#{key}"><span class="lv">Level<span class="n">{no}</span>{name}</span>'
                f'<span class="what"><b>{art}</b>{desc}<div class="meta">{meta}</div></span></a>')
HERO.append('</div>')
HERO.append('<div class="legendline"><span class="down">▼ cascade</span> = budgets and targets set at the level '
            'above &nbsp;·&nbsp; <span class="up">▲ ladder up</span> = every lower row names the higher row it '
            'moves (&ldquo;ladders to&rdquo;)</div>')
HERO.append('</div></div>')

FOOTER = f"""<footer><div class="in">
<h3>Behind this site</h3>
<ul>
<li>Editable masters: <span class="mono">kra_training_sheet.xlsx</span> (tracker + legend + CSI, FullStack &amp; CS Core and Content&ndash;Central team views) · <span class="mono">role_cards.xlsx</span> (ladder, 8 tabs)<span class="dlwrap" hidden> — download: {dlbtn("kra_training_sheet.xlsx")} {dlbtn("role_cards.xlsx")}</span> — in <span class="mono">docs/handoff/artifacts/</span> with the builders and CHANGELOG.</li>
<li>Standalone artifacts (updated in place; this site re-embeds them on republish): <a href="{LEVELS[0][6]}" target="_blank" rel="noopener">HOD one-pager</a> · <a href="{LEVELS[1][6]}" target="_blank" rel="noopener">KPI tracker</a> · <a href="{LEVELS[3][6]}" target="_blank" rel="noopener">AI Engineer Ladder</a>.</li>
<li>Comp bands are deliberately kept off this site — they live in the standalone ladder artifact and <span class="mono">role_cards.xlsx</span>.</li>
</ul>
</div></footer>"""

# ---- wide tables scroll in their own container (page never h-scrolls) -------
def wrap_tables(body):
    tags = re.findall(r"</?table", body)
    assert tags and len(tags) % 2 == 0
    assert all(t == ("<table" if i % 2 == 0 else "</table") for i, t in enumerate(tags))
    body = body.replace("<table", '<div class="tscroll"><table')
    return body.replace("</table>", "</table></div>")

for _c in (op_css, tr_css, ld_css, tm_css, SITE_CSS, op_body, tr_body, tm_body, ld_body):
    assert "tscroll" not in _c
op_body, tr_body, tm_body, ld_body = map(wrap_tables, (op_body, tr_body, tm_body, ld_body))

page = []
page.append("<title>Content OS</title>")
page.append("<style>")
page.append(SITE_CSS)
page.append(scope(op_css, "#org"))
page.append(scope(tr_css, "#dept"))
page.append(scope(ld_css, "#ind"))
page.append(f"#team {{\n{tm_css}\n}}")
page.append("#org, #dept, #team, #ind { padding: 40px 20px; }")
page.append(".tscroll { overflow-x: auto; }")
page.append("</style>")
page.append(NAV)
page.extend(HERO)
page.append(band(*LEVELS[0]))
page.append(f'<div class="site-section" id="org">\n{op_body}\n</div>')
page.append(band(*LEVELS[1]))
page.append(f'<div class="site-section" id="dept">\n{tr_body}\n</div>')
page.append(band(*LEVELS[2]))
page.append(f'<div class="site-section" id="team">\n{tm_body}\n</div>')
page.append(band(*LEVELS[3]))
page.append(f'<div class="site-section" id="ind">\n{ld_body}\n</div>')
page.append(FOOTER)
page.append(tr_scripts[0])
page.append(DL_SCRIPT)

out = "\n".join(page) + "\n"
assert "AI Engineer 3" not in out and "was Associate SDE" not in out
assert "What changed in v3" not in out
assert out.count('class="dlbtn"') == 5          # dept + team + ind bands, footer ×2
open("content_os.html", "w", encoding="utf-8").write(out)
print("chars:", len(out), "| braces:", out.count("{") == out.count("}"),
      "| sections:", [k for k, *_ in LEVELS], "| A chips:", len(A_ROWS),
      "| B rows:", len(B_ROWS), "| asks:", len(C_ROWS),
      "| dlbtns:", out.count('class="dlbtn"'),
      "| b64 chars:", {k: len(v) for k, v in XLSX_B64.items()})
